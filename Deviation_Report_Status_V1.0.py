import os
import xlwt,xlrd,xlsxwriter
import docx
import re
import msvcrt
from docx import Document
from xlutils.copy import copy
import tkinter as tk
import win32com.client
from win32com.client import Dispatch
from tkinter import filedialog
from tkinter import *
global doc_op_num, xls_op_num, pattern
pattern = '^[A-Z][0-9][0-9][0-9][0-9]$'
dm_pattern = '^.*([0-9][0-9][0-9][0-9][A-Z][0-9]+[A-Z][0-9]+)'
req_pattern_1 = '^.*([A-Z][A-Z][A-Z][0-9]+_[0-9]+)'
req_pattern_2 = '^.*([A-Z]-[A-Z][0-9][0-9][0-9][A-Z][A-Z][0-9]+_[0-9]+)'
doc_op_num = 1

#Ask progress the program
print("\n:::::::This Program is for SAIFEI V&V to Extract Deviation Report Data ::::::\n")
print("Contiune?(Y/N)")
in_content = input()
if in_content.lower() == "y":
    print("\nPlease Select Target Folder\n")
elif in_content.lower() == "n":
    exit(0)
else:
    print("Error input！")
    exit(0)

#Define addrerss list for docx and xlsx
docx_list = list()
xls_list = list()
global saveas_list
saveas_list = list()
error_list = list()
global folder_ad
folder_ad = filedialog.askdirectory()

#Create and Setup Workbook for Recording
summary_xls = folder_ad.replace("/","\\") + '\\Deviation_Report_Status.xlsx'
dev_ext = xlsxwriter.Workbook(summary_xls)
data_sheet = dev_ext.add_worksheet('Deviation Items Collection')
data_sheet.set_column('A:A',45)
data_sheet.set_column('B:D',25)
head_col = {'Deviation Report Name':0, 'Requirement ID':1, 'DM Number':2, 'Deviation Items':3}
for title in head_col:
    data_sheet.write(0,head_col[title],title)


#Save doc as docx
def docxize(address, document_name):
    old_path = address + "\\" + document_name
    new_path = address + "\\" + re.findall('(^.*)\.', document_name)[0] + '_copy.docx'
    w = win32com.client.Dispatch('Word.Application')
    w.Visible = 0
    w.DisplayAlerts = 0
    op_doc = w.Documents.Open(old_path)
    op_doc.SaveAs(new_path, 12, False, "", True, "", False, False, False, False)
    op_doc.Close()
    w.Quit()
    saveas_list.append(new_path)
    return new_path

#Read through entire documents in filedialog
ph_ad = os.walk(folder_ad)
for rt, fol, doc in ph_ad:
    new_ad = ''
    rt = rt.replace("/","\\")
    ad = rt
    if doc:
        for doc_name in doc:
            new_ad = ad + "\\" + doc_name
            if re.search("^DEV.*\.doc.*",doc_name):
                if re.search("^DEV.*\.docx",doc_name):
                    docx_list.append(new_ad)
                else:
                    docx_list.append(docxize(ad,doc_name))
            elif re.search("^\~\$.*",doc_name):
                continue
            else:
                error_list.append(doc_name)

print("Total " + str(len(docx_list)) + " Deviation Reports were Found.","\n")

#Write Record for Error
if error_list:
    warning_text = 'Connot Identified Following Documents:'
    error_file = folder_ad + '\\Error_Record.txt'
    txt = open(error_file,'w')
    txt.write(warning_text + "\n")
    for error_doc in error_list:
        txt.write(error_doc + "\n")
    txt.close()

#Extract data from Word documents
for doc in docx_list:
    itm_list = list()
    dc = Document(doc)
    name =  re.findall('^.*(DEV.*)\.', doc)[0]
    if re.search('^.*_copy', name):
        name = name.replace('_copy', '')
    if re.search(req_pattern_1, doc):
        doc_req = re.findall(req_pattern_1,doc)[0]
    elif re.search(req_pattern_2, doc):
        doc_req = re.findall(req_pattern_2,doc)[0]
    else:
        doc_req = 'Extract Error'
    if re.search(dm_pattern, doc):
        dm_num = re.findall(dm_pattern,doc)[0]
    else:
        dm_num = 'Extract Error'
    #ad = re.findall('(^.*)\\',doc)[0]
    tables = dc.tables
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                if re.search('^item.*',cell.text.lower()):
                    if cell.text not in itm_list:
                        itm_list.append(cell.text)
    data_sheet.write(doc_op_num, head_col['Deviation Report Name'], name)
    data_sheet.write(doc_op_num, head_col['Requirement ID'], doc_req)
    data_sheet.write(doc_op_num, head_col['DM Number'], dm_num)
    data_sheet.write(doc_op_num, head_col['Deviation Items'], len(itm_list))
    doc_op_num += 1


#Save and Terminate Program
dev_ext.close()
for new_doc in saveas_list:
    os.remove(new_doc)
print("\nWorks have been done!\n")

#Exit program
print("Press anykey to exit...")
while True:
    if ord(msvcrt.getch())!=None :
        break
